# export/excel.py
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from llm.extractor import SOURCE_TYPE_LABELS, OWNERSHIP_LABELS

# 导出排序：高价值来源类型靠前
_SOURCE_TYPE_ORDER = {
    "social": 0,
    "news_media": 1,
    "community": 2,
    "community_ugc": 3,
    "public_record": 4,
    "academic_technical": 5,
    "event": 6,
    "marketplace_directory": 7,
    "other": 8,
    "irrelevant": 99,
}

SOURCE_COLUMNS = [
    ("url", "URL"),
    ("domain", "域名"),
    ("source_type", "来源类型"),
    ("ownership", "归属"),
    ("confidence", "置信度"),
    ("bfs_round", "BFS轮次"),
    ("discovery_path", "发现路径"),
    ("title", "标题"),
    ("snippet", "搜索摘要"),
    ("deep_extracted", "深度抽取"),
    ("company_name_on_page", "页面公司名"),
    ("profile_or_handle", "账号/主页"),
    ("email", "邮箱"),
    ("phone", "电话"),
    ("evidence_snippet", "证据片段"),
    ("detail_confidence", "深度抽取置信度"),
    ("page_links_count", "页面链接数"),
    ("json_ld_types", "JSON-LD类型"),
    ("platform", "平台(名录)"),
    ("score", "检索得分"),
    ("weighted_score", "加权得分"),
]

PLATFORM_COLUMNS = [
    ("platform", "平台"),
    ("display_name", "展示名"),
    ("count", "链接数"),
    ("top_urls", "示例链接"),
]

CONDENSED_COLUMNS = [
    ("domain", "域名"),
    ("source_type", "来源类型"),
    ("ownership", "归属"),
    ("count", "该域链接数"),
    ("confidence", "代表链接置信度"),
    ("weighted_score", "代表链接加权分"),
    ("title", "代表标题"),
    ("url", "代表 URL"),
]


def _safe_filename(company: str) -> str:
    return re.sub(r"[^\w.-]+", "_", company).strip("_") or "company"


def _cell_value(value):
    """去掉 openpyxl 不允许的 XML 控制字符（常见于网页/LLM 抽取文本）。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _write_sheet(ws, headers: list[str], rows: list[list], freeze: bool = True) -> None:
    header_font = Font(bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, value in enumerate(row, 1):
            value = _cell_value(value)
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[c_idx - 1] == "URL" and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
    if freeze:
        ws.freeze_panes = "A2"
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        max_len = len(str(headers[col - 1]))
        for row in range(2, min(ws.max_row, 200) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, min(len(str(val)), 60))
        ws.column_dimensions[letter].width = max(10, min(max_len + 2, 50))


def _source_sort_key(source: dict) -> tuple:
    st = source.get("source_type") or "other"
    own = source.get("ownership") or "unknown"
    return (
        _SOURCE_TYPE_ORDER.get(st, 8),
        own != "first_party",
        -(source.get("weighted_score") or 0),
        -(source.get("confidence") or 0),
    )


def _source_rows(result: dict) -> list[list]:
    from search.denoise import platform_key

    rows = []
    for s in sorted(result.get("sources", []), key=_source_sort_key):
        detail = s.get("detail") or {}
        contacts = detail.get("contacts") or {}
        st = s.get("source_type") or "other"
        own = s.get("ownership") or "unknown"
        rows.append(
            [
                s.get("url", ""),
                s.get("domain", ""),
                SOURCE_TYPE_LABELS.get(st, st),
                OWNERSHIP_LABELS.get(own, own),
                s.get("confidence"),
                s.get("bfs_round", 0),
                s.get("discovery_path", ""),
                s.get("title", ""),
                s.get("snippet", ""),
                "是" if s.get("deep_extracted") else "否",
                detail.get("company_name_on_page", ""),
                detail.get("profile_or_handle", ""),
                contacts.get("email", ""),
                contacts.get("phone", ""),
                detail.get("evidence_snippet", ""),
                detail.get("is_same_company_confidence") if detail else None,
                detail.get("page_links_count"),
                ", ".join(detail.get("json_ld_types") or []) if detail else "",
                platform_key(s.get("url", "")) if st == "marketplace_directory" else "",
                s.get("score"),
                s.get("weighted_score"),
            ]
        )
    return rows


def _condensed_rows(result: dict) -> list[list]:
    """按域名聚合：每个域只展示一条代表链接 + 该域总数。"""
    from collections import defaultdict
    from search.denoise import host_of

    by_domain: dict[str, list[dict]] = defaultdict(list)
    for s in result.get("sources", []):
        by_domain[s.get("domain") or host_of(s.get("url", ""))].append(s)

    rows = []
    domain_items = sorted(
        by_domain.items(),
        key=lambda item: _source_sort_key(
            max(item[1], key=lambda s: (s.get("weighted_score") or 0, s.get("confidence") or 0))
        ),
    )
    for domain, sources in domain_items:
        best = max(sources, key=lambda s: (s.get("weighted_score") or 0, s.get("confidence") or 0))
        st = best.get("source_type") or "other"
        own = best.get("ownership") or "unknown"
        rows.append(
            [
                domain,
                SOURCE_TYPE_LABELS.get(st, st),
                OWNERSHIP_LABELS.get(own, own),
                len(sources),
                best.get("confidence"),
                best.get("weighted_score"),
                best.get("title", ""),
                best.get("url", ""),
            ]
        )
    return rows


def _summary_rows(result: dict) -> list[list]:
    anchor = result.get("anchor") or {}
    summary = result.get("summary") or {}
    rows = [
        ["公司", result.get("company", "")],
        ["官网", anchor.get("official_domain", "")],
        ["行业", anchor.get("industry", "")],
        ["生成时间", result.get("generated_at", "")],
        ["检索后端", result.get("search_backend", "")],
        ["可信来源总数", summary.get("total_sources", 0)],
        ["深度抽取数", summary.get("deep_extracted", 0)],
        [],
        ["来源类型", "数量"],
    ]
    for st, count in sorted(
        (summary.get("by_source_type") or {}).items(),
        key=lambda x: -x[1],
    ):
        rows.append([SOURCE_TYPE_LABELS.get(st, st), count])
    rows.append([])
    rows.append(["归属", "数量"])
    for own, count in (summary.get("by_ownership") or {}).items():
        rows.append([OWNERSHIP_LABELS.get(own, own), count])
    return rows


def _platform_rows(result: dict) -> list[list]:
    platforms = (result.get("summary") or {}).get("platforms") or []
    rows = []
    for p in platforms:
        tops = p.get("top_urls") or []
        rows.append(
            [
                p.get("platform", ""),
                p.get("display_name", ""),
                p.get("count", 0),
                "\n".join(tops),
            ]
        )
    return rows


def save_excel(result: dict, output_dir: str) -> Path:
    """把结构化结果写入 OUTPUT_DIR/<公司名>.xlsx。"""
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{_safe_filename(result.get('company', 'company'))}.xlsx"

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "摘要"
    _write_sheet(ws_summary, ["字段", "值"], _summary_rows(result), freeze=False)

    ws_sources = wb.create_sheet("足迹清单")
    _write_sheet(ws_sources, [h for _, h in SOURCE_COLUMNS], _source_rows(result))

    ws_condensed = wb.create_sheet("按域名汇总")
    _write_sheet(ws_condensed, [h for _, h in CONDENSED_COLUMNS], _condensed_rows(result))

    ws_platforms = wb.create_sheet("平台聚合")
    _write_sheet(ws_platforms, [h for _, h in PLATFORM_COLUMNS], _platform_rows(result))

    wb.save(path)
    return path
