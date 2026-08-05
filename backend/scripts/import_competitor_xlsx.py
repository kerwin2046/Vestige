"""Import peer/competitor Excel lists into Vestige as Targets (competitor role).

Sources:
  - 同行列表-2024.1.31 (2).xlsx  (Sheet1 + 常见)
  - 中国同行背调-2025-11-24.xlsx  (国内同行的基本情况)

Match key: normalized official domain.
Idempotent upsert (same rules as ExpoMind sync).

Usage:
  make import-competitors
  PYTHONPATH=backend python3 -m scripts.import_competitor_xlsx --dry-run
"""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from database import Database
from models import Company, utc_now
from scripts.sync_expomind import Incoming, merge_roles, normalize_domain, tier_rank


ROOT = Path(__file__).resolve().parents[2]

DEFAULT_FILES = [
    ROOT / "同行列表-2024.1.31 (2).xlsx",
    ROOT / "中国同行背调-2025-11-24.xlsx",
]


def _vestige_db_url() -> str:
    return os.getenv("VESTIGE_DATABASE_URL", "sqlite:///output/vestige.db")


def _clean_url(value: str | None) -> str:
    """Take first usable http(s) URL from a cell (may be multi-line)."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    for line in re.split(r"[\n\r]+", text):
        line = line.strip()
        if not line:
            continue
        # Skip social / non-company pages when a better URL may exist later
        low = line.lower()
        if "linkedin.com" in low or "facebook.com" in low:
            continue
        if "://" not in line and "." in line:
            line = f"https://{line}"
        if line.startswith("http://") or line.startswith("https://"):
            return line.split()[0].strip()
    # Fallback: first line even if social
    first = text.splitlines()[0].strip()
    if first and "://" not in first and "." in first:
        first = f"https://{first}"
    return first


def _clean_name(value: str | None, domain: str) -> str:
    name = (str(value).strip() if value is not None else "") or ""
    name = re.sub(r"\s+", " ", name)
    return name or domain


def _location_from_country(country: str | None, address: str | None = None) -> str:
    country = (str(country).strip() if country else "")
    if country:
        return country
    if address:
        # Address cells often end with country name
        lines = [ln.strip() for ln in str(address).splitlines() if ln.strip()]
        if lines:
            return lines[-1][:128]
    return ""


def load_peer_list(path: Path) -> list[Incoming]:
    items: list[Incoming] = []
    wb = load_workbook(path, read_only=True, data_only=True)

    if "Sheet1" in wb.sheetnames:
        ws = wb["Sheet1"]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c or "").strip().lower() for c in (rows[0] if rows else [])]
        # Expect Company, Web, Summary, Country, Address
        for row in rows[1:]:
            if not row or not any(row):
                continue
            cells = {header[i]: row[i] for i in range(min(len(header), len(row)))}
            url = _clean_url(cells.get("web") or cells.get("website") or cells.get("url"))
            domain = normalize_domain(url)
            if not domain:
                continue
            name = _clean_name(cells.get("company") or cells.get("name"), domain)
            summary = str(cells.get("summary") or "").strip()
            items.append(
                Incoming(
                    name=name,
                    domain=domain,
                    industry="additive / manufacturing peers",
                    location=_location_from_country(
                        cells.get("country"), cells.get("address")
                    ),
                    tier="target",
                    roles=["competitor"],
                    priority="",
                    provenance={
                        "origin": "xlsx.peer_list.sheet1",
                        "file": path.name,
                        "sheet": "Sheet1",
                        "url": url,
                        "summary": summary[:500] if summary else None,
                    },
                )
            )

    if "常见" in wb.sheetnames:
        ws = wb["常见"]
        rows = list(ws.iter_rows(values_only=True))
        for row in rows[1:]:
            if not row or not any(row):
                continue
            name_raw, link = (row[0] if len(row) > 0 else None), (
                row[1] if len(row) > 1 else None
            )
            url = _clean_url(link)
            domain = normalize_domain(url)
            if not domain:
                continue
            name = _clean_name(name_raw, domain)
            items.append(
                Incoming(
                    name=name,
                    domain=domain,
                    industry="additive / manufacturing peers",
                    location="",
                    tier="target",
                    roles=["competitor"],
                    priority="High",
                    provenance={
                        "origin": "xlsx.peer_list.common",
                        "file": path.name,
                        "sheet": "常见",
                        "url": url,
                    },
                )
            )

    wb.close()
    return items


def load_china_peers(path: Path) -> list[Incoming]:
    items: list[Incoming] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = "国内同行的基本情况"
    if sheet not in wb.sheetnames:
        wb.close()
        return items

    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    for row in rows[1:]:
        if not row or not any(row):
            continue
        name_raw = row[0] if len(row) > 0 else None
        link = row[1] if len(row) > 1 else None
        url = _clean_url(link)
        domain = normalize_domain(url)
        if not domain:
            continue
        # Capture alternate domains from multi-line cells
        extra_domains: list[str] = []
        for line in re.split(r"[\n\r]+", str(link or "")):
            d = normalize_domain(_clean_url(line))
            if d and d != domain and d not in extra_domains:
                if "linkedin.com" in d:
                    continue
                extra_domains.append(d)

        name = _clean_name(name_raw, domain)
        items.append(
            Incoming(
                name=name,
                domain=domain,
                industry="中国快速成型 / CNC 同行",
                location="China",
                tier="target",
                roles=["competitor"],
                priority="High",
                provenance={
                    "origin": "xlsx.china_peers",
                    "file": path.name,
                    "sheet": sheet,
                    "url": url,
                    "extra_domains": extra_domains or None,
                },
            )
        )
    wb.close()
    return items


def upsert_competitor(session, item: Incoming, *, dry_run: bool, source: str) -> str:
    existing = session.scalar(
        select(Company).where(Company.official_domain == item.domain)
    )
    if existing:
        changed = False
        if tier_rank(item.tier) > tier_rank(existing.tier or "candidate"):
            existing.tier = item.tier
            changed = True
        new_roles = merge_roles(existing.roles, item.roles)
        if new_roles != (existing.roles or []):
            existing.roles = new_roles
            changed = True
        if item.industry and not existing.industry:
            existing.industry = item.industry
            changed = True
        if item.location and not existing.location:
            existing.location = item.location
            changed = True
        if item.priority and (
            not existing.priority
            or (item.priority == "High" and existing.priority != "High")
        ):
            existing.priority = item.priority
            changed = True
        if item.name and item.name != existing.name:
            aliases = list(existing.aliases or [])
            if item.name not in aliases and item.name != existing.name:
                aliases.append(item.name)
                existing.aliases = aliases
                changed = True
        # Prefer High-priority Chinese / 常见 display names when existing looks generic
        if (
            item.priority == "High"
            and item.name
            and existing.name
            and existing.name.casefold() != item.name.casefold()
            and len(item.name) >= len(existing.name)
        ):
            aliases = list(existing.aliases or [])
            if existing.name not in aliases:
                aliases.append(existing.name)
            existing.aliases = aliases
            existing.name = item.name
            changed = True
        prov = dict(existing.provenance or {})
        prov.update({k: v for k, v in item.provenance.items() if v is not None})
        sources = sorted(
            set(
                (prov.get("sources") or [])
                + [item.provenance.get("origin", source), source]
            )
        )
        prov["sources"] = sources
        if prov != (existing.provenance or {}):
            existing.provenance = prov
            changed = True
        if existing.source in {"", "manual", "expomind"} and source:
            # Keep stronger labeling without wiping expomind if already there
            if existing.source != source:
                existing.source = (
                    f"{existing.source}+{source}"
                    if existing.source and existing.source != "manual"
                    else source
                )
                changed = True
        if changed:
            existing.updated_at = utc_now()
            return "updated"
        return "unchanged"

    if dry_run:
        return "would_create"

    session.add(
        Company(
            name=item.name,
            official_domain=item.domain,
            industry=item.industry,
            location=item.location,
            aliases=[],
            tier=item.tier,
            roles=item.roles,
            priority=item.priority,
            source=source,
            provenance={
                **{k: v for k, v in item.provenance.items() if v is not None},
                "sources": [item.provenance.get("origin", source), source],
            },
        )
    )
    return "created"


def prefer_item(prev: Incoming | None, item: Incoming) -> Incoming:
    if prev is None:
        return item
    # Higher priority wins; High > empty
    prev_score = 2 if prev.priority == "High" else 1
    item_score = 2 if item.priority == "High" else 1
    if item_score > prev_score:
        item.roles = merge_roles(prev.roles, item.roles)
        item.provenance = {**prev.provenance, **item.provenance}
        if prev.location and not item.location:
            item.location = prev.location
        if prev.industry and not item.industry:
            item.industry = prev.industry
        # Keep alternate name as alias via provenance
        if prev.name and prev.name != item.name:
            item.provenance["aka"] = sorted(
                set((item.provenance.get("aka") or []) + [prev.name])
            )
        return item
    prev.roles = merge_roles(prev.roles, item.roles)
    prev.provenance = {**item.provenance, **prev.provenance}
    if item.location and not prev.location:
        prev.location = item.location
    if item.name and item.name != prev.name:
        prev.provenance["aka"] = sorted(
            set((prev.provenance.get("aka") or []) + [item.name])
        )
    return prev


def run_import(paths: list[Path], *, dry_run: bool = False) -> dict:
    incoming: list[Incoming] = []
    for path in paths:
        if not path.exists():
            raise FileNotFoundError(path)
        name = path.name
        if "中国同行" in name or "背调" in name:
            incoming.extend(load_china_peers(path))
        else:
            incoming.extend(load_peer_list(path))

    by_domain: dict[str, Incoming] = {}
    for item in incoming:
        by_domain[item.domain] = prefer_item(by_domain.get(item.domain), item)

    database = Database(_vestige_db_url())
    database.create_all()
    stats = {
        "created": 0,
        "updated": 0,
        "unchanged": 0,
        "would_create": 0,
        "incoming": len(by_domain),
        "files": [p.name for p in paths],
    }

    with database.session_factory() as session:
        for item in by_domain.values():
            result = upsert_competitor(
                session, item, dry_run=dry_run, source="competitor-xlsx"
            )
            stats[result] = stats.get(result, 0) + 1
        if not dry_run:
            session.commit()
    database.dispose()
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import peer Excel lists as Vestige competitor Targets"
    )
    parser.add_argument(
        "files",
        nargs="*",
        type=Path,
        help="Excel paths (default: both peer list files in project root)",
    )
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    paths = list(args.files) if args.files else list(DEFAULT_FILES)
    stats = run_import(paths, dry_run=args.dry_run)
    print(json.dumps({"dry_run": args.dry_run, **stats}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
