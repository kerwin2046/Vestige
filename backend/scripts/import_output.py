"""Import historical discovery results from output/*.xlsx and output/*.json into SQLite.

Usage (from repo root, venv active):
  PYTHONPATH=backend python3 -m scripts.import_output
  make import-output
"""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from openpyxl import load_workbook
from sqlalchemy import select

from database import Database
from llm.extractor import OWNERSHIP_LABELS, SOURCE_TYPE_LABELS
from models import Company, Run, RunStatus, utc_now
from repositories import runs as runs_repo

SOURCE_TYPE_REVERSE = {label: key for key, label in SOURCE_TYPE_LABELS.items()}
OWNERSHIP_REVERSE = {label: key for key, label in OWNERSHIP_LABELS.items()}

# Some older exports used slightly different wording.
SOURCE_TYPE_REVERSE.update(
    {
        "社区/论坛/Q&A/评论": "community",
        "社媒": "social",
        "其他": "other",
        "招聘": "recruitment",
        "展会/会议": "event",
        "自有/官方发布物": "owned",
        "参考/数据库(维基、企业信息库、行业百科)": "reference",
    }
)

SKIP_NAMES = {"test.xlsx", "vestige.db"}


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def _output_dir() -> Path:
    return Path(os.getenv("VESTIGE_OUTPUT_DIR", _repo_root() / "output"))


def _database_url() -> str:
    default = os.getenv("VESTIGE_DATABASE_URL")
    if default:
        return default
    # Keep relative path so it matches make api / worker defaults when cwd is repo root.
    return "sqlite:///output/vestige.db"


def _host_of(url: str) -> str:
    try:
        host = (urlsplit(url).hostname or "").lower()
    except Exception:
        return ""
    return host.removeprefix("www.")


def _normalize_domain(value: str | None) -> str:
    value = (value or "").strip().lower()
    if not value:
        return ""
    if "://" not in value:
        value = f"//{value}"
    host = urlsplit(value).hostname or ""
    return host.lower().removeprefix("www.")


def _as_float(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _as_int(value: Any, default: int = 0) -> int:
    if value is None or value == "":
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _map_source_type(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "other"
    if raw in SOURCE_TYPE_LABELS:
        return raw
    return SOURCE_TYPE_REVERSE.get(raw, raw if re.fullmatch(r"[a-z0-9_]+", raw) else "other")


def _map_ownership(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "unknown"
    if raw in OWNERSHIP_LABELS:
        return raw
    return OWNERSHIP_REVERSE.get(raw, "unknown")


def _read_summary_meta(ws) -> dict[str, Any]:
    meta: dict[str, Any] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        if key == "公司":
            meta["name"] = str(value or "").strip()
        elif key == "官网":
            meta["official_domain"] = _normalize_domain(str(value or ""))
        elif key == "行业":
            meta["industry"] = str(value or "").strip()
        elif key == "地区" or key == "位置":
            meta["location"] = str(value or "").strip()
        elif key == "检索后端":
            meta["search_backend"] = str(value or "").strip()
        elif key == "生成时间":
            meta["generated_at"] = str(value or "").strip()
    return meta


def _row_dict(headers: tuple[Any, ...], values: tuple[Any, ...]) -> dict[str, Any]:
    return {
        str(h).strip(): values[i] if i < len(values) else None
        for i, h in enumerate(headers)
        if h is not None and str(h).strip()
    }


def _source_from_row(row: dict[str, Any]) -> dict[str, Any] | None:
    url = str(row.get("URL") or row.get("url") or "").strip()
    if not url.startswith("http"):
        return None
    domain = str(row.get("域名") or row.get("domain") or "").strip() or _host_of(url)
    detail: dict[str, Any] = {}
    company_name = row.get("页面公司名")
    profile = row.get("账号/主页")
    evidence = row.get("证据片段")
    email = row.get("邮箱")
    phone = row.get("电话")
    if any([company_name, profile, evidence, email, phone]):
        detail = {
            "company_name_on_page": str(company_name or ""),
            "profile_or_handle": str(profile or ""),
            "evidence_snippet": str(evidence or ""),
            "contacts": {
                "email": str(email or ""),
                "phone": str(phone or ""),
            },
        }
        conf = row.get("深度抽取置信度")
        if conf not in (None, ""):
            detail["is_same_company_confidence"] = _as_float(conf)

    return {
        "url": url,
        "canonical_url": url,
        "domain": domain,
        "source_type": _map_source_type(row.get("来源类型") or row.get("source_type")),
        "ownership": _map_ownership(row.get("归属") or row.get("ownership")),
        "confidence": _as_float(row.get("置信度") or row.get("confidence"), 0.0),
        "title": str(row.get("标题") or row.get("title") or ""),
        "snippet": str(row.get("搜索摘要") or row.get("snippet") or ""),
        "discovery_path": str(row.get("发现路径") or row.get("discovery_path") or "imported"),
        "bfs_round": _as_int(row.get("BFS轮次") or row.get("bfs_round"), 0),
        "detail": detail or None,
    }


def parse_excel(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "摘要" not in wb.sheetnames or "足迹清单" not in wb.sheetnames:
            raise ValueError(f"{path.name}: missing 摘要/足迹清单 sheets")
        meta = _read_summary_meta(wb["摘要"])
        if not meta.get("name"):
            meta["name"] = path.stem
        rows = list(wb["足迹清单"].iter_rows(values_only=True))
        if not rows:
            return meta, []
        headers = rows[0]
        sources: list[dict[str, Any]] = []
        seen: set[str] = set()
        for values in rows[1:]:
            item = _source_from_row(_row_dict(headers, values))
            if not item:
                continue
            key = item["url"]
            if key in seen:
                continue
            seen.add(key)
            sources.append(item)
        return meta, sources
    finally:
        wb.close()


def parse_json(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    anchor = data.get("anchor") or {}
    meta = {
        "name": str(data.get("company") or anchor.get("name") or path.stem).strip(),
        "official_domain": _normalize_domain(
            anchor.get("official_domain") or data.get("official_domain") or ""
        ),
        "industry": str(anchor.get("industry") or ""),
        "location": str(anchor.get("location") or ""),
        "search_backend": str(data.get("search_backend") or ""),
        "generated_at": str(data.get("generated_at") or ""),
    }
    sources: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in data.get("sources") or []:
        url = str(item.get("url") or "").strip()
        if not url.startswith("http") or url in seen:
            continue
        seen.add(url)
        sources.append(
            {
                "url": url,
                "canonical_url": item.get("canonical_url") or url,
                "domain": item.get("domain") or _host_of(url),
                "source_type": _map_source_type(item.get("source_type")),
                "ownership": _map_ownership(item.get("ownership")),
                "confidence": _as_float(item.get("confidence"), 0.0),
                "title": item.get("title") or "",
                "snippet": item.get("snippet") or "",
                "discovery_path": item.get("discovery_path")
                or ",".join(item.get("matched_queries") or [])
                or "imported",
                "bfs_round": _as_int(item.get("bfs_round"), 0),
                "detail": item.get("detail"),
            }
        )
    return meta, sources


def _find_company(session, *, name: str, official_domain: str) -> Company | None:
    companies = list(session.scalars(select(Company)))
    if official_domain:
        for company in companies:
            if (company.official_domain or "").lower() == official_domain:
                return company
    needle = name.casefold()
    for company in companies:
        if company.name.casefold() == needle:
            return company
    return None


def _get_or_create_company(
    session,
    *,
    name: str,
    official_domain: str = "",
    industry: str = "",
    location: str = "",
) -> Company:
    company = _find_company(session, name=name, official_domain=official_domain)
    if company is None:
        company = Company(
            name=name,
            official_domain=official_domain,
            industry=industry,
            location=location,
            aliases=[],
        )
        session.add(company)
        session.commit()
        session.refresh(company)
        return company

    changed = False
    if official_domain and not company.official_domain:
        company.official_domain = official_domain
        changed = True
    if industry and not company.industry:
        company.industry = industry
        changed = True
    if location and not company.location:
        company.location = location
        changed = True
    if changed:
        company.updated_at = utc_now()
        session.commit()
        session.refresh(company)
    return company


def _already_imported(session, *, company_id: str, export_path: str) -> bool:
    existing = list(
        session.scalars(
            select(Run).where(
                Run.company_id == company_id,
                Run.export_path == export_path,
            )
        )
    )
    return bool(existing)


def import_file(session, path: Path, *, dry_run: bool = False) -> dict[str, Any]:
    rel = str(path.relative_to(_repo_root())) if path.is_absolute() else str(path)
    if path.suffix.lower() == ".json":
        meta, sources = parse_json(path)
    else:
        meta, sources = parse_excel(path)

    name = meta.get("name") or path.stem
    if dry_run:
        return {
            "file": rel,
            "company": name,
            "sources": len(sources),
            "action": "dry-run",
        }

    company = _get_or_create_company(
        session,
        name=name,
        official_domain=meta.get("official_domain") or "",
        industry=meta.get("industry") or "",
        location=meta.get("location") or "",
    )
    if _already_imported(session, company_id=company.id, export_path=rel):
        return {
            "file": rel,
            "company": company.name,
            "company_id": company.id,
            "sources": len(sources),
            "action": "skipped",
        }

    run = Run(
        company_id=company.id,
        status=RunStatus.SUCCEEDED,
        stage="succeeded",
        progress=100,
        settings_snapshot={
            "imported": True,
            "search_backend": meta.get("search_backend") or None,
            "generated_at": meta.get("generated_at") or None,
            "source_file": rel,
        },
        export_path=rel,
        started_at=utc_now(),
        finished_at=utc_now(),
    )
    session.add(run)
    session.commit()
    session.refresh(run)

    runs_repo.replace_run_sources(session, run.id, sources)
    runs_repo.append_run_event(
        session,
        run_id=run.id,
        stage="imported",
        message=f"Imported {len(sources)} sources from {path.name}",
        payload={"source_count": len(sources), "file": rel},
    )
    return {
        "file": rel,
        "company": company.name,
        "company_id": company.id,
        "run_id": run.id,
        "sources": len(sources),
        "action": "imported",
    }


def discover_files(output_dir: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(output_dir.iterdir(), key=lambda p: p.name.casefold()):
        if not path.is_file():
            continue
        if path.name in SKIP_NAMES or path.name.startswith("."):
            continue
        if path.suffix.lower() not in {".xlsx", ".json"}:
            continue
        files.append(path)

    # Prefer JSON over Excel for the same stem (richer payload).
    by_stem: dict[str, list[Path]] = {}
    for path in files:
        by_stem.setdefault(path.stem.casefold(), []).append(path)
    chosen: list[Path] = []
    for group in by_stem.values():
        jsons = [p for p in group if p.suffix.lower() == ".json"]
        chosen.extend(jsons or group)
    return sorted(chosen, key=lambda p: p.name.casefold())


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Import output Excel/JSON into Vestige SQLite")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=_output_dir(),
        help="Directory containing historical exports (default: ./output)",
    )
    parser.add_argument("--dry-run", action="store_true", help="Parse only, do not write DB")
    args = parser.parse_args(argv)

    output_dir = args.output_dir.resolve()
    if not output_dir.exists():
        print(f"output dir not found: {output_dir}")
        return 1

    files = discover_files(output_dir)
    if not files:
        print(f"no importable files in {output_dir}")
        return 0

    database = Database(_database_url())
    database.create_all()

    imported = skipped = failed = 0
    total_sources = 0
    print(f"db={_database_url()}")
    print(f"files={len(files)} dry_run={args.dry_run}")

    with database.session_factory() as session:
        for path in files:
            try:
                result = import_file(session, path, dry_run=args.dry_run)
            except Exception as exc:
                failed += 1
                print(f"FAIL {path.name}: {type(exc).__name__}: {exc}")
                continue
            total_sources += int(result.get("sources") or 0)
            action = result.get("action")
            if action == "imported":
                imported += 1
            elif action == "skipped":
                skipped += 1
            print(
                f"{action:8} {result.get('company')} · {result.get('sources')} sources · {path.name}"
            )

    print(
        f"done imported={imported} skipped={skipped} failed={failed} source_rows={total_sources}"
    )
    return 0 if failed == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
